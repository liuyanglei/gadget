"""Collect complete Chinese financial disclosures and official documents."""

from __future__ import annotations

import hashlib
import io
import json
import logging
import re
import shutil
import subprocess
import tempfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "data" / "news.json"
SETTINGS_FILE = ROOT / "data" / "settings.json"
DEFAULT_CATEGORY_LIMIT = 500
SSE_DISCOVERY_LIMIT = 400
MAX_DOWNLOAD_BYTES = 25 * 1024 * 1024
MAX_PDF_PAGES = 80
MAX_OCR_PAGES = 8
OCR_TIMEOUT_SECONDS = 25
MIN_CONTENT_LENGTH = 180
TIMEOUT_SECONDS = 30
WORKERS = 6
CATEGORIES = ("A股公告", "分红派息", "公司报告", "监管动态", "宏观数据", "政策解读")
CHINA_TZ = timezone(timedelta(hours=8))

MIIT_LISTING_URL = "https://wap.miit.gov.cn/RRSdy/index.html"
MIIT_ORIGIN = "https://www.miit.gov.cn"
SSE_API = "https://query.sse.com.cn/security/stock/queryCompanyBulletin.do"
SSE_ORIGIN = "https://www.sse.com.cn"
SSE_FILE_ORIGIN = "https://static.sse.com.cn"
HEADERS = {
    "User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/140.0.0.0 Safari/537.36",
    "Accept-Language": "zh-CN,zh;q=0.9",
}

logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
logging.getLogger("pypdf").setLevel(logging.ERROR)


def normalized_text(value: str) -> str:
    lines = [re.sub(r"[ \t\u3000]+", " ", line).strip() for line in value.splitlines()]
    return "\n".join(line for line in lines if line)


def article_id(value: str) -> str:
    return hashlib.sha1(value.encode("utf-8")).hexdigest()[:16]


def iso_datetime(value: str) -> str:
    value = value.strip()
    for pattern in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(value, pattern).replace(tzinfo=CHINA_TZ).isoformat()
        except ValueError:
            continue
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")


def category_for(title: str, source_kind: str) -> str:
    if re.search(r"权益分派|利润分配|现金分红|现金红利|派息|分红派息|股息", title):
        return "分红派息"
    if re.search(r"年度报告|半年度报告|季度报告|业绩说明|业绩预告|业绩快报|投资者关系|调研活动|经营数据|研究报告|评估报告|审计报告|募集说明书|上市保荐书|法律意见书", title):
        return "公司报告"
    if source_kind == "sse":
        return "A股公告"
    if re.search(r"经济运行|行业运行|统计公报|统计数据|增加值|产量|销量|增长|月份.*情况|上半年.*情况|前\d+个月", title):
        return "宏观数据"
    if re.search(r"解读|意见|办法|规定|规范|政策|指南", title):
        return "政策解读"
    return "监管动态"


def download(session: requests.Session, url: str, referer: str | None = None) -> bytes:
    headers = {"Referer": referer} if referer else None
    response = session.get(url, timeout=TIMEOUT_SECONDS, headers=headers, stream=True)
    response.raise_for_status()
    if "sse.com.cn" in url and response.headers.get("Content-Type", "").startswith("text/html"):
        challenge = re.search(r"arg1='([0-9A-F]+)'", response.text)
        if challenge:
            positions = [15, 35, 29, 24, 33, 16, 1, 38, 10, 9, 19, 31, 40, 27, 22, 23, 25, 13, 6, 11,
                         39, 18, 20, 8, 14, 21, 32, 26, 2, 30, 7, 4, 17, 5, 3, 28, 34, 37, 12, 36]
            shuffled = [""] * 40
            for index, char in enumerate(challenge.group(1)):
                for output_index, position in enumerate(positions):
                    if position == index + 1:
                        shuffled[output_index] = char
            rearranged = "".join(shuffled)
            mask = "3000176000856006061501533003690027800375"
            cookie = "".join(
                f"{int(rearranged[index:index + 2], 16) ^ int(mask[index:index + 2], 16):02x}"
                for index in range(0, min(len(rearranged), len(mask)), 2)
            )
            session.cookies.set("acw_sc__v2", cookie, domain=".sse.com.cn", path="/")
            response.close()
            response = session.get(url, timeout=TIMEOUT_SECONDS, headers=headers, stream=True)
            response.raise_for_status()
    length = int(response.headers.get("Content-Length") or 0)
    if length > MAX_DOWNLOAD_BYTES:
        raise ValueError(f"attachment exceeds {MAX_DOWNLOAD_BYTES} bytes")
    chunks: list[bytes] = []
    size = 0
    for chunk in response.iter_content(1024 * 256):
        size += len(chunk)
        if size > MAX_DOWNLOAD_BYTES:
            raise ValueError(f"attachment exceeds {MAX_DOWNLOAD_BYTES} bytes")
        chunks.append(chunk)
    return b"".join(chunks)


def image_ocr(image) -> str:
    try:
        import pytesseract
        return normalized_text(pytesseract.image_to_string(image, lang="chi_sim+eng", timeout=OCR_TIMEOUT_SECONDS))
    except Exception:
        return ""


def extract_pdf(data: bytes) -> tuple[str, dict[str, int]] | None:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            return None
        page_count = len(reader.pages)
        if not page_count or page_count > MAX_PDF_PAGES:
            return None
        texts: list[str] = []
        needs_ocr: list[int] = []
        for index, page in enumerate(reader.pages):
            text = normalized_text(page.extract_text() or "")
            if len(text) < 20 and len(page.images) > 0:
                needs_ocr.append(index + 1)
            texts.append(text)

        ocr_pages = 0
        if needs_ocr:
            if len(needs_ocr) > MAX_OCR_PAGES:
                return None
            try:
                from pdf2image import convert_from_bytes
                for page_number in needs_ocr:
                    images = convert_from_bytes(data, dpi=220, first_page=page_number, last_page=page_number)
                    ocr_text = image_ocr(images[0]) if images else ""
                    if len(ocr_text) < 20:
                        return None
                    texts[page_number - 1] = ocr_text
                    ocr_pages += 1
            except Exception as exc:  # OCR is optional locally but mandatory for image-only pages.
                logging.warning("OCR failed: %s", exc)
                return None

        content = normalized_text("\n\n".join(texts))
        if len(content) < MIN_CONTENT_LENGTH:
            return None
        return content, {"pages": page_count, "ocr_pages": ocr_pages}
    except Exception as exc:
        logging.warning("PDF extraction failed: %s", exc)
        return None


def extract_docx(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        parts.extend("\t".join(cell.text.strip() for cell in row.cells) for row in table.rows)
    return normalized_text("\n".join(parts))


def extract_xlsx(data: bytes) -> str:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in workbook.worksheets:
        parts.append(f"【{sheet.title}】")
        for row in sheet.iter_rows(values_only=True):
            values = [str(value).strip() if value is not None else "" for value in row]
            if any(values):
                parts.append("\t".join(values))
    return normalized_text("\n".join(parts))


def extract_legacy_office(data: bytes, suffix: str) -> str:
    command = shutil.which("antiword") if suffix == ".doc" else shutil.which("libreoffice")
    if not command:
        return ""
    with tempfile.TemporaryDirectory() as temp_dir:
        source = Path(temp_dir) / f"attachment{suffix}"
        source.write_bytes(data)
        if suffix == ".doc":
            result = subprocess.run([command, str(source)], capture_output=True, timeout=90)
            return normalized_text(result.stdout.decode("utf-8", errors="ignore"))
        subprocess.run([command, "--headless", "--convert-to", "xlsx", "--outdir", temp_dir, str(source)], capture_output=True, timeout=120)
        converted = source.with_suffix(".xlsx")
        return extract_xlsx(converted.read_bytes()) if converted.exists() else ""


def extract_attachment(data: bytes, url: str) -> tuple[str, dict[str, int]] | None:
    suffix = Path(urlparse(url).path).suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(data)
    if suffix == ".docx":
        text = extract_docx(data)
    elif suffix == ".xlsx":
        text = extract_xlsx(data)
    elif suffix in {".doc", ".xls"}:
        text = extract_legacy_office(data, suffix)
    elif suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}:
        from PIL import Image
        text = image_ocr(Image.open(io.BytesIO(data)))
    else:
        return None
    text = normalized_text(text)
    return (text, {"pages": 0, "ocr_pages": 1 if suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"} else 0}) if len(text) >= MIN_CONTENT_LENGTH else None


def discover_sse(session: requests.Session, start_date: str, end_date: str) -> list[dict[str, str]]:
    params = {
        "isPagination": "true",
        "securityType": "0101,120100,020100,020200,120200",
        "reportType": "ALL",
        "beginDate": start_date,
        "endDate": end_date,
        "pageHelp.pageSize": str(SSE_DISCOVERY_LIMIT),
        "pageHelp.pageNo": "1",
        "pageHelp.beginPage": "1",
        "pageHelp.cacheSize": "1",
        "pageHelp.endPage": "5",
    }
    response = session.get(SSE_API, params=params, headers={"Referer": SSE_ORIGIN}, timeout=TIMEOUT_SECONDS)
    response.raise_for_status()
    rows = response.json().get("pageHelp", {}).get("data", [])
    results = []
    for row in rows[:SSE_DISCOVERY_LIMIT]:
        url = urljoin(SSE_FILE_ORIGIN, row.get("URL") or "")
        title = normalized_text(row.get("TITLE") or "")
        if title and url.endswith(".pdf"):
            results.append({
                "id": article_id(url),
                "title": title,
                "url": url,
                "published_at": iso_datetime(row.get("ADDDATE") or row.get("SSEDATE") or ""),
                "source": f"上海证券交易所 · {row.get('SECURITY_CODE', '')} {row.get('SECURITY_NAME', '')}".strip(),
                "category": category_for(title, "sse"),
                "source_kind": "sse",
            })
    return results


def extract_sse(entry: dict[str, str]) -> dict[str, object] | None:
    session = requests.Session()
    session.headers.update(HEADERS)
    try:
        data = download(session, entry["url"], SSE_ORIGIN)
        if not data.startswith(b"%PDF-"):
            raise ValueError("download did not return a PDF")
        result = extract_pdf(data)
        if not result:
            return None
        content, stats = result
        return {key: value for key, value in entry.items() if key not in {"url", "source_kind"}} | {"content": content, "attachment_stats": stats}
    except (requests.RequestException, ValueError) as exc:
        logging.warning("Skipping SSE attachment %s: %s", entry["url"], exc)
        return None


def discover_miit(session: requests.Session) -> list[str]:
    response = session.get(MIIT_LISTING_URL, timeout=TIMEOUT_SECONDS)
    response.raise_for_status()
    soup = BeautifulSoup(response.content, "html.parser")
    urls: list[str] = []
    seen: set[str] = set()
    for anchor in soup.select('a[href*="/art/"]'):
        href = str(anchor.get("href") or "")
        url = urljoin(MIIT_ORIGIN, href)
        if href.endswith(".html") and url not in seen:
            seen.add(url)
            urls.append(url)
    return urls[:45]


def extract_miit(session: requests.Session, url: str) -> dict[str, object] | None:
    response = session.get(url, timeout=TIMEOUT_SECONDS)
    response.raise_for_status()
    soup = BeautifulSoup(response.content, "html.parser")
    content_node = soup.select_one("#con_con")
    title_node = soup.select_one("#con_title")
    if not content_node or not title_node:
        return None
    title = normalized_text(title_node.get_text(" ", strip=True))
    parts = [normalized_text(content_node.get_text("\n", strip=True))]
    attachment_count = 0
    for anchor in content_node.select("a[href]"):
        attachment_url = urljoin(url, str(anchor.get("href") or ""))
        suffix = Path(urlparse(attachment_url).path).suffix.lower()
        if suffix not in {".pdf", ".doc", ".docx", ".xls", ".xlsx", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}:
            continue
        try:
            result = extract_attachment(download(session, attachment_url, url), attachment_url)
            if not result:
                return None
            attachment_text, _ = result
            parts.append(f"附件：{normalized_text(anchor.get_text(' ', strip=True))}\n{attachment_text}")
            attachment_count += 1
        except (requests.RequestException, ValueError):
            return None
    content = normalized_text("\n\n".join(parts))
    if len(content) < MIN_CONTENT_LENGTH:
        return None
    date_meta = soup.select_one('meta[name="PubDate"]')
    source_meta = soup.select_one('meta[name="ContentSource"]')
    published = iso_datetime(str(date_meta.get("content") if date_meta else ""))
    detail = normalized_text(str(source_meta.get("content") if source_meta else ""))
    return {
        "id": article_id(url),
        "title": title,
        "published_at": published,
        "content": content,
        "source": f"工业和信息化部 · {detail}" if detail else "工业和信息化部",
        "category": category_for(title, "miit"),
        "attachment_stats": {"attachments": attachment_count},
    }


def load_previous() -> list[dict[str, object]]:
    try:
        payload = json.loads(OUTPUT.read_text(encoding="utf-8"))
        articles = [item for item in payload.get("articles", []) if item.get("content")]
        for item in articles:
            source_kind = "sse" if str(item.get("source", "")).startswith("上海证券交易所") else "miit"
            item["category"] = category_for(str(item.get("title", "")), source_kind)
            published_at = str(item.get("published_at", ""))
            if published_at.endswith("Z"):
                legacy_time = datetime.fromisoformat(published_at.removesuffix("Z"))
                item["published_at"] = legacy_time.replace(tzinfo=CHINA_TZ).isoformat()
        return [item for item in articles if is_publishable(item)]
    except (OSError, json.JSONDecodeError, AttributeError):
        return []


def is_publishable(item: dict[str, object]) -> bool:
    title = normalized_text(str(item.get("title", "")))
    content = normalized_text(str(item.get("content", "")))
    category = str(item.get("category", ""))
    source = str(item.get("source", ""))
    if not title or len(content) < MIN_CONTENT_LENGTH or category not in CATEGORIES:
        return False
    if source.startswith("上海证券交易所"):
        stats = item.get("attachment_stats") or {}
        return isinstance(stats, dict) and int(stats.get("pages") or 0) > 0
    if source.startswith("工业和信息化部"):
        finance_terms = re.compile(
            r"企业|行业|经济|工业|制造|汽车|新能源|通信|软件|电池|金属|市场|投资|产业|生产|消费|经营|标准|政策|税|回款"
        )
        return bool(finance_terms.search(title + "\n" + content[:1000]))
    return False


def load_category_limits() -> dict[str, int]:
    limits = {category: DEFAULT_CATEGORY_LIMIT for category in CATEGORIES}
    try:
        settings = json.loads(SETTINGS_FILE.read_text(encoding="utf-8"))
        configured = settings.get("category_limits", {})
        for category in CATEGORIES:
            value = int(configured.get(category, limits[category]))
            limits[category] = min(5000, max(20, value))
    except (OSError, ValueError, TypeError, json.JSONDecodeError, AttributeError):
        logging.warning("Invalid settings.json; using safe defaults")
    return limits


def main() -> None:
    session = requests.Session()
    session.headers.update(HEADERS)
    category_limits = load_category_limits()
    previous = load_previous()
    previous_ids = {str(item.get("id")) for item in previous}
    today = datetime.now(CHINA_TZ).date()
    start_date = (today - timedelta(days=1)).isoformat()
    end_date = today.isoformat()

    collected: list[dict[str, object]] = []
    try:
        sse_entries = discover_sse(session, start_date, end_date)
        new_sse = [item for item in sse_entries if item["id"] not in previous_ids]
        logging.info("Discovered %d SSE notices (%d new) for %s to %s", len(sse_entries), len(new_sse), start_date, end_date)
        with ThreadPoolExecutor(max_workers=WORKERS) as executor:
            futures = [executor.submit(extract_sse, item) for item in new_sse]
            for future in as_completed(futures):
                article = future.result()
                if article:
                    collected.append(article)
    except (requests.RequestException, ValueError) as exc:
        logging.warning("SSE discovery failed: %s", exc)

    try:
        for url in discover_miit(session):
            candidate_id = article_id(url)
            if candidate_id in previous_ids:
                continue
            try:
                article = extract_miit(session, url)
                if article:
                    collected.append(article)
            except requests.RequestException as exc:
                logging.warning("Skipping MIIT page %s: %s", url, exc)
    except requests.RequestException as exc:
        logging.warning("MIIT discovery failed: %s", exc)

    combined = collected + previous
    unique: dict[str, dict[str, object]] = {}
    for item in combined:
        key = re.sub(r"\W+", "", str(item.get("title", "")).lower())
        if key and is_publishable(item) and key not in unique:
            unique[key] = item
    sorted_articles = sorted(unique.values(), key=lambda item: str(item.get("published_at", "")), reverse=True)
    category_counts = {category: 0 for category in CATEGORIES}
    articles: list[dict[str, object]] = []
    for item in sorted_articles:
        category = str(item["category"])
        if category_counts[category] >= category_limits[category]:
            continue
        articles.append(item)
        category_counts[category] += 1
    payload = {
        "updated_at": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
        "successful_sources": len({str(item.get("source", "")).split(" · ")[0] for item in articles}),
        "content_mode": "official_full_text_zh_with_attachments",
        "backfill_days": 2,
        "category_limits": category_limits,
        "articles": articles,
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    logging.info("Wrote %d complete articles (%d newly extracted)", len(articles), len(collected))


if __name__ == "__main__":
    main()
